"""Safe local discovery and source-preserving chunking. No network access."""
from __future__ import annotations

import ast
import hashlib
import logging
import os
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from pathspec import GitIgnoreSpec

TEXT_SUFFIXES = {
    ".py", ".pyi", ".js", ".jsx", ".ts", ".tsx", ".rs", ".go", ".c", ".h", ".cpp", ".hpp",
    ".java", ".kt", ".swift", ".rb", ".php", ".sh", ".zsh", ".bash", ".sql", ".r", ".m",
    ".md", ".mdx", ".txt", ".rst", ".toml", ".yaml", ".yml", ".json", ".xml", ".html", ".css",
    ".scss", ".tex", ".csv", ".ini", ".cfg", ".vue", ".svelte",
}
SKIP_DIRS = {"node_modules", "vendor", "dist", "build", "target", "__pycache__", "venv", "env"}
# Excluded before opening. Explicitly never inspect credential stores.
SECRET_NAME = re.compile(r"(^\.env($|\.)|credentials|secrets?|id_rsa|id_ed25519|\.pem$|\.key$|\.p12$|\.pfx$)", re.I)
SECRET_VALUE = re.compile(
    r"-----BEGIN (?:[A-Z ]*PRIVATE KEY)-----|\b(?:gh[pousr]_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{30,}|"
    r"AKIA[A-Z0-9]{16}|sk-(?:proj-)?[A-Za-z0-9_-]{24,})\b|"
    r"(?i:(?:api[_-]?key|password|access[_-]?token|client[_-]?secret)\s*[:=]\s*[\"'][^\"'\s]{12,}[\"'])"
)
MAX_FILE_BYTES = 2 * 1024 * 1024


@dataclass
class Source:
    path: str
    digest: str
    text: str
    kind: str


@dataclass
class Chunk:
    start: int
    end: int
    text: str
    symbol: str = ""
    calls: tuple[tuple[int, str], ...] = ()


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def tokenize(text: str) -> list[str]:
    expanded = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", text)
    return re.findall(r"[^\W_]+", expanded.replace("_", " ").lower(), re.UNICODE)


def extract(path: Path, data: bytes) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF support requires ravel-index[documents]") from exc
        import io
        from pypdf.errors import PyPdfError
        # pypdf warning text can contain untrusted header bytes. Drop its library
        # logging while parsing; restore the prior setting even on parser errors.
        logger = logging.getLogger("pypdf")
        previous_handlers, previous_propagate = logger.handlers[:], logger.propagate
        logger.handlers = [logging.NullHandler()]
        logger.propagate = False
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 500:
                raise ValueError("Encrypted or oversized PDF excluded")
            return "\n".join(p.extract_text() or "" for p in reader.pages), "pdf-extracted"
        except (PyPdfError, RecursionError, KeyError, TypeError) as exc:
            raise ValueError("Unreadable PDF excluded") from exc
        finally:
            logger.handlers = previous_handlers
            logger.propagate = previous_propagate
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX support requires ravel-index[documents]") from exc
        import io
        from docx.opc.exceptions import OpcError
        from lxml.etree import XMLSyntaxError
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError("Oversized DOCX excluded")
            doc = Document(io.BytesIO(data))
            return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells)
                            for table in doc.tables for row in table.rows]), "docx-extracted"
        except (zipfile.BadZipFile, OpcError, XMLSyntaxError, KeyError) as exc:
            raise ValueError("Unreadable DOCX excluded") from exc
    if b"\x00" in data:
        raise ValueError("Binary file excluded")
    return data.decode("utf-8"), "source"


def discover(root: Path):
    """Yield eligible files, honoring hierarchical gitignore and .ravelignore rules.

    Permission errors are not retried through alternative access paths. Symlinks
    and hidden directories are excluded. Tracked files also honor current ignores.
    """
    def walk(directory: Path, inherited: list[tuple[Path, GitIgnoreSpec]]):
        rules = list(inherited)
        for name in (".gitignore", ".ravelignore"):
            ignore = directory / name
            try:
                if not ignore.is_symlink():
                    rules.append((directory, GitIgnoreSpec.from_lines(ignore.read_text().splitlines())))
            except (OSError, UnicodeError):
                pass
        try:
            entries = sorted(os.scandir(directory), key=lambda e: e.name)
        except OSError:
            return
        for entry in entries:
            path = Path(entry.path)
            try:
                if entry.is_symlink() or entry.name.startswith(".") or SECRET_NAME.search(entry.name):
                    continue
                is_dir = entry.is_dir(follow_symlinks=False)
                if is_dir and entry.name in SKIP_DIRS:
                    continue
                ignored = False
                for base, spec in rules:
                    check = path.relative_to(base).as_posix() + ("/" if is_dir else "")
                    match = spec.check_file(check)
                    if match.include is not None:
                        ignored = match.include
                if ignored:
                    continue
                if is_dir:
                    yield from walk(path, rules)
                elif entry.is_file(follow_symlinks=False):
                    if path.suffix.lower() in TEXT_SUFFIXES | {".pdf", ".docx"} or entry.name in {
                        "Dockerfile", "Makefile", "LICENSE", "NOTICE"}:
                        yield path
            except OSError:
                continue
    # Respect ignore files above the explicitly selected root as well.
    parents = []
    for parent in reversed(root.parents):
        for name in (".gitignore", ".ravelignore"):
            try:
                p = parent / name
                if not p.is_symlink():
                    parents.append((parent, GitIgnoreSpec.from_lines(p.read_text().splitlines())))
            except (OSError, UnicodeError):
                pass
    yield from walk(root, parents)


def read_source(root: Path, path: Path) -> Source | None:
    try:
        # O_NOFOLLOW closes the file-symlink race on platforms supporting it.
        fd = os.open(path, os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0))
        with os.fdopen(fd, "rb") as f:
            stat = os.fstat(f.fileno())
            if stat.st_size > MAX_FILE_BYTES or stat.st_mode & 0o444 == 0:
                return None
            data = f.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            return None
        text, kind = extract(path, data)
        if SECRET_VALUE.search(text):
            return None
        return Source(path.relative_to(root).as_posix(), digest(data), text, kind)
    except (OSError, UnicodeError, ValueError):
        return None


def chunks(text: str, suffix: str = "", width: int = 70) -> list[Chunk]:
    lines = text.splitlines()
    if not lines:
        return []
    spans: list[tuple[int, int, str, tuple[str, ...]]] = []
    call_lines = []
    if suffix == ".py":
        try:
            tree = ast.parse(text)
            call_lines = [(n.lineno, ast.unparse(n.func)) for n in ast.walk(tree) if isinstance(n, ast.Call)]
            def visit(node, prefix=""):
                for child in ast.iter_child_nodes(node):
                    if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                        name = prefix + child.name
                        start = min([child.lineno] + [d.lineno for d in child.decorator_list])
                        calls = tuple(sorted({ast.unparse(n.func) for n in ast.walk(child)
                                             if isinstance(n, ast.Call)}))
                        spans.append((start, child.end_lineno or start, name, calls))
                        visit(child, name + ".")
                    else:
                        visit(child, prefix)
            visit(tree)
        except (SyntaxError, RecursionError):
            pass
    # Non-overlapping source slices, with the innermost containing Python symbol.
    boundaries = {1, len(lines) + 1}
    for start, end, _, _ in spans:
        boundaries.update((start, end + 1))
    points = sorted(boundaries)
    result = []
    for a, b in zip(points, points[1:]):
        containing = sorted((s for s in spans if s[0] <= a and s[1] >= b - 1), key=lambda s: s[1] - s[0])
        symbol, calls = (containing[0][2], containing[0][3]) if containing else ("", ())
        for start in range(a, b, width):
            end = min(start + width, b) - 1
            body = "\n".join(lines[start - 1:end])
            if body.strip():
                result.append(Chunk(start, end, body, symbol, tuple((line, call) for line, call in call_lines
                                                                      if start <= line <= end)))
    return result
